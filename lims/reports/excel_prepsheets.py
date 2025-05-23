import sys
import os

# Get the absolute path to the root "LIMS" directory
current_file = os.path.abspath(__file__)
lims_root = os.path.abspath(os.path.join(current_file, "../../.."))

# Insert it at the start of sys.path
sys.path.insert(0, lims_root)

from lims.config.config import CONNECTION_STRING
import pandas as pd
import lims.config.tables as tables
import lims.config.lab_lists as lab_lists
from lims.reports.batch_summary import GenerateBatchSummary
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from lims.config import file_paths
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_from_string
from openpyxl.utils import column_index_from_string, get_column_letter

class GenerateExcelPrepsheets:
    @staticmethod
    def __init__():
        session = None
        engine = None

    def init_session():
        # Initialize the SQLAlchemy session
        engine = create_engine(CONNECTION_STRING)
        tables.Base.metadata.create_all(engine)
        Session = sessionmaker(bind=engine)
        session = Session()
        return session

    def generate_prepsheets(sdg):
        try:
            session = GenerateExcelPrepsheets.init_session()

            query = session.query(tables.DQO).filter(tables.DQO.SDG == sdg).all()

            # Convert list of SQLAlchemy model instances to list of dicts
            query_data = [row.__dict__ for row in query]

            # Remove SQLAlchemy internal state, if present
            for row in query_data:
                row.pop('_sa_instance_state', None)

            # Create the DataFrame
            dqo = pd.DataFrame(query_data)

            print(dqo)

            batch_id_list = dqo['BatchID'].unique().tolist()

            print(batch_id_list)

        except Exception as e:
            print(f"An exception occurred: {e}")
        finally:
            if session:
                session.close()

        try:
            session = GenerateExcelPrepsheets.init_session()

            query = session.query(tables.SampleLogin).filter(tables.SampleLogin.SDG == sdg).all()

            # Convert list of SQLAlchemy model instances to list of dicts
            query_data = [row.__dict__ for row in query]

            # Remove SQLAlchemy internal state, if present
            for row in query_data:
                row.pop('_sa_instance_state', None)

            # Create the DataFrame
            batch_summary = pd.DataFrame(query_data)

            column_order = [
                "SDG",
                "SampleID",
                "Matrix",
                "HG",
                "ISOAM",
                "ISOPU",
                "ISOTH",
                "ISOU",
                "GAMMA",
                "GFPC",
                "LSCAB",
                "LSCPU",
                "LSCSR",
                "MET",
                "TCLP",
                "BEF",
                "SIO2",
                'FLUOR',
                "NH3",
                "NO3",
                "NO2",
                "CRVI",
                "CL",
                "PH",
                "TSS",
                "TSP",
                "SampleDate",
                "SampleTime",
                "DateReceived",
                "TimeReceived"
            ]

            batch_summary = batch_summary.reindex(columns=column_order)

            batch_summary.replace(True, "x", inplace=True)
            batch_summary.replace(False, "", inplace=True)

        except Exception as e:
            print(f"An exception occurred: {e}")
        finally:
            if session:
                session.close()

        try:
            session = GenerateExcelPrepsheets.init_session()

            query = session.query(tables.CoC).filter(tables.CoC.SDG == sdg).first()

            coc = query.CoCID

        except Exception as e:
            print(f"An exception occurred: {e}")
        finally:
            if session:
                session.close()

        GenerateBatchSummary.generate_batch_summary(coc, batch_summary)

        for batch in batch_id_list:
            GenerateExcelPrepsheets.fill_excel_template(batch, dqo, batch_summary)

    def fill_excel_template(batch_id, dqo, batch_summary):
        method = dqo[dqo['BatchID']==batch_id]['Method'].unique().tolist()[0]
        matrix = dqo[dqo['BatchID']==batch_id]['Matrix'].unique().tolist()[0]
        sdg = dqo["SDG"].unique().tolist()[0]

        print(method)

        if method in lab_lists.matrix_dependent_templates:
            template_file = f"{method} ({matrix}).xlsx"
        else:
            template_file = f"{method}.xlsx"

        if template_file in lab_lists.excel_template_field_locations.keys():
            field_locations = lab_lists.excel_template_field_locations[template_file]
        else:
            return
        
        template_file_path = os.path.join(file_paths.prepsheet_template_directory, template_file)
        print(template_file_path)
        
        # Load the Excel workbook and active sheet
        wb = load_workbook(template_file_path)
        ws = wb.active  # or use wb[sheet_name] if your templates have named sheets

        # Fill the BatchID field if present
        if 'BatchID' in field_locations:
            cell_address = field_locations['BatchID']
            ws[cell_address] = batch_id
        if field_locations.get("Matrix"):
            cell_address = field_locations['Matrix']
            ws[cell_address] = matrix

        batch_view = dqo[dqo["BatchID"] == batch_id].copy()

        # Make result type column
        batch_view['ResultType'] = batch_view['SampleID'].apply(GenerateExcelPrepsheets.determine_result_type)

        result_type_order = lab_lists.excel_result_type_order

        # Convert 'ResultType' to a categorical type with the specified order
        batch_view['ResultType'] = pd.Categorical(batch_view['ResultType'], categories=result_type_order, ordered=True)

        # Sort the DataFrame by this custom order
        batch_view = batch_view.sort_values('ResultType')

        #Find the DUP sample row if it exists
        dup_row = batch_view[batch_view['ResultType'] == 'DUP']

        if not dup_row.empty:
            dup_sample_id = dup_row['SampleID'].iloc[0]
            parent_sample_id = dup_sample_id.replace('DUP', '')

            print(f"DUP Sample: {dup_sample_id} -> Original Sample: {parent_sample_id}")

            # Get the integer position of the DUP sample in batch_view
            dup_pos = batch_view.index.get_loc(dup_row.index[0])  # converts label to position

            # Get the parent row, if it exists
            parent_row = batch_view[batch_view['SampleID'] == parent_sample_id]
            
            if not parent_row.empty:
                parent_index = parent_row.index[0]

                # Drop the parent row
                batch_view_reordered = batch_view.drop(index=parent_index)

                # Slice using iloc (integer positions)
                top = batch_view_reordered.iloc[:dup_pos + 1]
                bottom = batch_view_reordered.iloc[dup_pos + 1:]

                # Concatenate
                batch_view = pd.concat([top, parent_row, bottom]).reset_index(drop=True)
            else:
                print(f"Original sample {parent_sample_id} not found in batch.")
        else:
            print("No DUP sample found in this batch.")

        samples = batch_view['SampleID'].tolist()
        sample_cell_address = field_locations['SampleID']

        parent_samples = batch_view[batch_view['ResultType'] == 'REG']['SampleID'].unique().tolist()

        col_letter, sample_row_number = coordinate_from_string(sample_cell_address)
        col_index = column_index_from_string(col_letter)

        for sample in samples:
            # Write the SampleID
            ws[f"{get_column_letter(col_index)}{sample_row_number}"] = sample

            # Find the matching parent sample (if any)
            matching_parent = next((parent for parent in parent_samples if parent in sample), None)

            if matching_parent:
                # Filter the summary row for the parent sample
                summary_row = batch_summary[batch_summary['SampleID'] == matching_parent]

                if not summary_row.empty:
                    # Extract date and time from the filtered row
                    sample_date = summary_row.iloc[0]['SampleDate']
                    sample_time = summary_row.iloc[0]['SampleTime']

                    # Write the SampleDate in the next column
                    ws[f"{get_column_letter(col_index + 1)}{sample_row_number}"] = sample_date

                    # Write the SampleTime in the column after that
                    ws[f"{get_column_letter(col_index + 2)}{sample_row_number}"] = sample_time

            sample_row_number += 1

        sdg_directory = os.path.join(file_paths.sdg_directory, sdg)
        os.makedirs(sdg_directory, exist_ok=True)

        # Save the filled template (optionally use a new filename)
        prepsheet_destination_path = GenerateExcelPrepsheets.get_unique_filename(sdg_directory, f"{batch_id}-Prep")
        wb.save(prepsheet_destination_path)

    def get_unique_filename(directory, base_filename, extension=".xlsx"):
        """
        Check for existing filenames and apply Windows-style -1, -2 rule.
        
        Parameters:
            directory (str): The directory where the file will be saved.
            base_filename (str): The base filename without extension.
            extension (str): File extension (default is ".xlsx").
            
        Returns:
            str: A full path to a unique filename in the directory.
        """
        full_path = os.path.join(directory, f"{base_filename}{extension}")
        counter = 1

        while os.path.exists(full_path):
            full_path = os.path.join(directory, f"{base_filename}-{counter}{extension}")
            counter += 1

        return full_path

    def determine_result_type(sample_id):
        for qc in lab_lists.all_qc:
            if qc in sample_id:
                return qc
        return 'REG' 