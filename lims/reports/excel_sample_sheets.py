import sys
import os
import pandas as pd
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker
from openpyxl import load_workbook
from openpyxl.utils.cell import (
    coordinate_from_string, column_index_from_string,
    get_column_letter
)
from openpyxl.cell.cell import MergedCell

from lims.config.config import CONNECTION_STRING
import lims.config.tables as tables
import lims.config.lab_lists as lab_lists
from lims.config import file_paths


class GenerateSampleSheets:

    @staticmethod
    def init_session():
        engine = create_engine(CONNECTION_STRING)
        tables.Base.metadata.create_all(engine)
        Session = sessionmaker(bind=engine)
        return Session()

    @staticmethod
    def debugger(title, text):
        from PyQt5.QtWidgets import QMessageBox
        msg = QMessageBox()
        msg.setIcon(QMessageBox.Critical)
        msg.setWindowTitle(str(title))
        msg.setText(str(text))
        msg.setStandardButtons(QMessageBox.Ok)
        msg.exec_()

    @staticmethod
    def generate_sheets(sdg):
        """Generate prepsheet or preservation sheet depending on matrix type."""

        session = None
        try:
            session = GenerateSampleSheets.init_session()

            query = (
                session.query(
                    tables.SampleLogin.SDG,
                    tables.SampleLogin.Matrix,
                    tables.SampleLogin.LocationID,
                    tables.SampleLogin.SampleID
                )
                .filter(tables.SampleLogin.SDG == sdg)
                .all()
            )

            query_data = [dict(row._mapping) for row in query]
            df = pd.DataFrame(query_data)

        except Exception as e:
            print(f"An exception occurred: {e}")
            return
        finally:
            if session:
                session.close()

        if df.empty:
            GenerateSampleSheets.debugger("Sample Sheet Error", f"No samples found for SDG {sdg}")
            return

        matrix_list = {str(m).upper() for m in df['Matrix'].unique().tolist()}

        if len(matrix_list) != 1:
            GenerateSampleSheets.debugger(
                "Sample Sheet Error",
                f"SDG {sdg} contains more than one matrix type: {matrix_list}"
            )
            return

        matrix = matrix_list.pop()

        sheet_name = (
            "Sample Prepsheet" if matrix == "SO"
            else "Sample Preservation Sheet" if matrix == "AQ"
            else None
        )

        if sheet_name is None:
            GenerateSampleSheets.debugger("Sample Sheet Error", f"No sheet template mapped for matrix type: {matrix}")
            return

        GenerateSampleSheets.fill_sample_sheet(sheet_name, df, sdg)

    @staticmethod
    def fill_sample_sheet(sheet_name, merged_df, sdg):
        """Fill Excel sheet, creating new files if sample area is exceeded."""

        field_locations = lab_lists.sample_sheet_field_locations[sheet_name]

        template_file = sheet_name + ".xlsx"
        template_file_path = os.path.join(file_paths.prepsheet_template_directory, template_file)

        # Required for row boundary logic
        start_col_letter, start_row = coordinate_from_string(field_locations['SampleID'])
        start_col = column_index_from_string(start_col_letter)
        _, last_row = coordinate_from_string(field_locations['LastCell'])

        # Output directory for this SDG
        sdg_directory = os.path.join(file_paths.sdg_directory, sdg)
        os.makedirs(sdg_directory, exist_ok=True)

        file_counter = 0

        def load_template():
            wb = load_workbook(template_file_path)
            ws = wb.active
            ws[field_locations['SDG']] = sdg
            return wb, ws

        wb, ws = load_template()
        sample_row = start_row

        def save_sheet():
            nonlocal wb, ws, file_counter, sample_row
            file_counter += 1
            filename = f"{sdg}-{sheet_name}" if file_counter == 1 else f"{sdg}-{sheet_name}-{file_counter}"
            destination = GenerateSampleSheets.get_unique_filename(sdg_directory, filename)
            wb.save(destination)
            wb, ws = load_template()  # reset sheet for next batch
            sample_row = start_row     # reset row pointer

        def write_to_cell(cell_address, value):
            cell = ws[cell_address]
            if isinstance(cell, MergedCell):
                for merged_range in ws.merged_cells.ranges:
                    if cell.coordinate in merged_range:
                        tl = merged_range.min_row, merged_range.min_col
                        cell_address = f"{get_column_letter(tl[1])}{tl[0]}"
                        break
            ws[cell_address].value = value

        for _, sample in merged_df.iterrows():
            if sample_row > last_row:
                save_sheet()

            sample_id = sample['SampleID']
            field_id = sample.get('LocationID', '')

            write_to_cell(f"{get_column_letter(start_col)}{sample_row}", sample_id)
            write_to_cell(f"{get_column_letter(start_col + 1)}{sample_row}", field_id)

            sample_row += 1

        # Save last populated sheet
        save_sheet()

    @staticmethod
    def get_unique_filename(directory, base_filename, extension=".xlsx"):
        full_path = os.path.join(directory, f"{base_filename}{extension}")
        counter = 1

        while os.path.exists(full_path):
            full_path = os.path.join(directory, f"{base_filename}-{counter}{extension}")
            counter += 1

        return full_path
